import os
import glob
import cv2
import pandas as pd
import time
from tqdm import tqdm
from ultralytics import YOLO
from openpyxl import load_workbook


def append_row_to_excel(excel_path, row_dict):
    """append one result row to an excel workbook, creating the file if needed."""
    if not os.path.exists(excel_path):
        pd.DataFrame([row_dict]).to_excel(excel_path, index=False)
    else:
        wb = load_workbook(excel_path)
        ws = wb.active
        # initialize headers when appending to an empty worksheet
        if ws.max_row == 1 and ws.max_column == 1 and ws.cell(1, 1).value is None:
            ws.append(list(row_dict.keys()))
        ws.append([row_dict[k] for k in row_dict.keys()])
        wb.save(excel_path)


def process_videos(video_dir, output_dir, model_path="yolo11n.pt"):
    os.makedirs(output_dir, exist_ok=True)
    excel_path = os.path.join(output_dir, "video_detection_times.xlsx")

    video_files = glob.glob(os.path.join(video_dir, "*.mp4")) + \
                  glob.glob(os.path.join(video_dir, "*.MP4"))
    if not video_files:
        raise ValueError(f"no .mp4 files found in '{video_dir}'; check the input path")

    model = YOLO(model_path)
    computer_labels = {"laptop", "tv"}

    for video_path in tqdm(video_files, desc="processing videos", unit="video", dynamic_ncols=True):
        video_name = os.path.basename(video_path)
        stem = os.path.splitext(video_name)[0]
        bbox_dir = os.path.join(output_dir, stem)
        os.makedirs(bbox_dir, exist_ok=True)

        cap = cv2.VideoCapture(video_path)
        if not cap.isOpened():
            tqdm.write(f"error opening {video_name}; skipping")
            continue

        total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        fps = cap.get(cv2.CAP_PROP_FPS) or 0.0
        duration = total_frames / fps if fps else 0.0

        counts = {"person": 0, "computer": 0, "both": 0}
        frame_times = []

        t0_video = time.perf_counter()

        bar = tqdm(total=total_frames,
                   desc=f"frames {stem}",
                   unit="frame",
                   leave=False,
                   dynamic_ncols=True)

        frame_idx = 0
        while True:
            t0_read = time.perf_counter()
            ret, frame = cap.read()
            t1_read = time.perf_counter()

            if not ret:
                break

            try:
                t0_proc = time.perf_counter()
                results = model(frame)[0]
                t1_proc = time.perf_counter()
            except Exception as e:
                tqdm.write(f"{repr(bar)} frame {frame_idx} error: {e}")
                frame_times.append((t1_read - t0_read))
                frame_idx += 1
                bar.update(1)
                continue

            # save annotated frames for detections retained by the model
            if results.boxes is not None and len(results.boxes) > 0:
                img = results.plot()

                # downscale annotated frames to reduce storage costs
                h, w = img.shape[:2]
                small = cv2.resize(img, (w // 2, h // 2), interpolation=cv2.INTER_LINEAR)

                # write compressed jpeg review frames for later inspection
                fname = f"{stem}_frame_{frame_idx:06d}.jpg"
                cv2.imwrite(
                    os.path.join(bbox_dir, fname),
                    small,
                    [int(cv2.IMWRITE_JPEG_QUALITY), 25]
                )

                labels = [model.names[int(c)] for c in results.boxes.cls.cpu().numpy()]
                has_p = "person" in labels
                has_c = any(l in computer_labels for l in labels)
                if has_p and has_c:
                    counts["both"] += 1
                elif has_p:
                    counts["person"] += 1
                elif has_c:
                    counts["computer"] += 1

            frame_times.append((t1_read - t0_read) + (t1_proc - t0_proc))
            frame_idx += 1
            bar.update(1)

        bar.close()
        cap.release()

        t1_video = time.perf_counter()
        total_time = t1_video - t0_video
        avg_time_per_frame = (sum(frame_times) / len(frame_times)) if frame_times else 0.0
        throughput_fps = len(frame_times) / total_time if total_time else 0.0

        row = {
            "video_file": video_name,
            "num_frames": total_frames,
            "fps": fps,
            "duration_s": duration,
            "total_time_s": total_time,
            "avg_time_per_frame_s": avg_time_per_frame,
            "throughput_fps": throughput_fps,
            "time_person_s": counts["person"] / fps if fps else 0.0,
            "time_computer_s": counts["computer"] / fps if fps else 0.0,
            "time_both_s": counts["both"] / fps if fps else 0.0,
            "pct_frames_person": counts["person"] / total_frames if total_frames else 0.0,
            "pct_frames_computer": counts["computer"] / total_frames if total_frames else 0.0,
            "pct_frames_both": counts["both"] / total_frames if total_frames else 0.0,
        }

        append_row_to_excel(excel_path, row)
        tqdm.write(f"{repr(bar)} saved results for {video_name}")

    print(f"\nall done; incremental results saved to {excel_path}")


if __name__ == "__main__":
    import argparse

    p = argparse.ArgumentParser()
    p.add_argument("video_dir", help="directory containing .mp4 files")
    p.add_argument("output_dir", help="directory for annotated frames and excel output")
    p.add_argument("--model", default="yolo11n.pt", help="path to yolo model weights")
    args = p.parse_args()
    process_videos(args.video_dir, args.output_dir, args.model)
